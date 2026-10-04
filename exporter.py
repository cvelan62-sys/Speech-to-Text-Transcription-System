"""Export helpers: TXT and DOCX."""
import io
from datetime import datetime

from docx import Document
from docx.shared import Pt


def to_txt(text: str) -> bytes:
    return text.encode("utf-8")


def to_docx(text: str, source: str, language: str) -> bytes:
    doc = Document()
    doc.add_heading("Speech-to-Text Transcription", level=0)
    meta = doc.add_paragraph(f"Source: {source}   |   Language: {language}   |   Generated: {datetime.now():%d %b %Y, %H:%M}")
    meta.runs[0].font.size = Pt(9)
    for para in text.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
